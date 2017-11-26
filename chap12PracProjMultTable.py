#!python3
# 12multiplicationTable.py - Creates a multiplication table.

import openpyxl, sys
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# Store N from command line call.
if len(sys.argv) > 1:
    N = int(sys.argv[1])
else:
    print('Something went terribly wrong at the very beginning!')

# Create bold, frozen pane headers in row 1 & column 1.
wb = openpyxl.Workbook()
sheet = wb.active
sheet.freeze_panes = 'B2'
bold = Font(bold=True)
for i in range(2, N+2):
    sheet['A' + str(i)].font = bold
    sheet['A' + str(i)].value = i-1
for i in range(2, N+2):
    sheet[get_column_letter(i) + '1'].font = bold
    sheet[get_column_letter(i) + '1'].value = i-1

# Populate multiplication table
for i in range(2, N+2):
    for j in range(2, N+2):
        sheet[get_column_letter(i) + str(j)] = (i-1)*(j-1)

# Save file
wb.save('multTable' + str(N) + '.xlsx')
